import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_word_document():
    doc = Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("ATW WEB FLOW: 4D WEB FLOW EXPLORER")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(14, 116, 144) # Deep Cyan

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("The Secret Life of the Web Explained for Kids Under 12 in 4D Space-Time\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Embed Mascot Image
    img_path = os.path.abspath("assets/mascot.jpg")
    if os.path.exists(img_path):
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.add_run().add_picture(img_path, width=Inches(2.4))
        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cap = caption_p.add_run("Meet Professor Byte: Your 4D Web Guide!")
        run_cap.font.size = Pt(10)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph() # Spacing

    # Section 1: Executive Overview
    h1 = doc.add_heading("1. What is ATW WEB FLOW & Why '4D'?", level=1)
    h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    p1 = doc.add_paragraph(
        "Traditional web diagrams show the internet in flat, boring 2D lines. ATW WEB FLOW introduces "
        "Space (3D) AND Time (The 4th Dimension) so kids can truly understand what happens inside their devices "
        "in the split second between clicking a link and watching a video."
    )
    p1.paragraph_format.line_spacing = 1.25

    # 4D Highlights Bullet List
    b1 = doc.add_paragraph(style='List Bullet')
    r1 = b1.add_run("⏱️ 4D Temporal Time Scrubber: ")
    r1.bold = True
    b1.add_run("Kids can scrub time forwards and backwards between 0ms and 12ms to follow every millisecond of a web packet's flight.")

    b2 = doc.add_paragraph(style='List Bullet')
    r2 = b2.add_run("❄️ Freeze Time (Bullet-Time): ")
    r2.bold = True
    b2.add_run("Pause data packets in mid-air over the ocean to inspect their contents.")

    b3 = doc.add_paragraph(style='List Bullet')
    r3 = b3.add_run("👻 4D Ghost Echoes: ")
    r3.bold = True
    b3.add_run("Visualizes past and future trajectories simultaneously using holographic temporal shadows.")

    doc.add_paragraph()

    # Section 2: Educational Stations & Kid Analogies
    h2 = doc.add_heading("2. The 7 Educational Stations & Kid Analogies", level=1)
    h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    # Table of Stations
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    headers = ["Station", "3D Three.js Visual", "Kid-Friendly Analogy"]
    widths = [Inches(1.8), Inches(2.2), Inches(2.8)]

    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "0F172A") # Dark Navy

    stations = [
        ("1. The Client (You!) 💻", "Floating tech island with laptop, dynamic screen canvas & pet bot.", "You are the detective! Your phone or tablet shouts into space: 'Hey internet, show me cute cat videos!'"),
        ("2. DNS Satellite 🔮", "Asteroid pod with golden solar wings & scanning radar dish.", "The Magic Phonebook Owl: Turns names like roblox.com into computer coordinate numbers (IPs) in 2 milliseconds!"),
        ("3. Rainbow Fiber Highway 🌈", "Curved translucent tube with streaming photon particles.", "The Light Tunnel: Glass threads thinner than your hair shooting laser pulses under oceans at light speed!"),
        ("4. Cloud Server Citadel ☁️", "Futuristic skyscraper city with blinking LEDs & cooling fans.", "The Master Kitchen: A super-computer that never sleeps! It cooks up your HTML/CSS/JS files like a master chef."),
        ("5. API Waiter Drone 🤖", "Quadcopter hover drone carrying an animated JSON cargo box.", "The Restaurant Waiter: You don't cook in the kitchen—the waiter fetches fresh pizza! The API fetches live weather and scores."),
        ("6. Build-A-Page Lab 🧪", "Stage with interactive 3D robot: skeleton, armor & energy core.", "The Superhero Robot Trio: HTML is the bones, CSS is the painted armor/cape, and JS is the batteries and brain!"),
        ("7. Webpage vs Website 📖", "Holographic 3D magic book with animated turning pages.", "The Comic Book: A Webpage is ONE single page. A Website is the entire comic book bound by hyperlinks!")
    ]

    for idx, (st_name, st_vis, st_analogy) in enumerate(stations):
        row_cells = table.add_row().cells
        bg_col = "F8FAFC" if idx % 2 == 0 else "EDF2F7"
        for i, val in enumerate([st_name, st_vis, st_analogy]):
            row_cells[i].width = widths[i]
            p = row_cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            set_cell_background(row_cells[i], bg_col)

    doc.add_paragraph()

    # Section 3: The Superhero Trio (HTML, CSS, JS)
    h3 = doc.add_heading("3. The Superhero Web Trio: HTML, CSS & JavaScript", level=1)
    h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    p3 = doc.add_paragraph(
        "Inside the 3D 'Build-A-Page Lab', kids can flip toggle switches on and off to witness how the "
        "three foundational languages of the web interact:"
    )

    trio_items = [
        ("🦴 HTML (HyperText Markup Language): ", "The Skeleton & Lego Frame. It creates the raw boxes, paragraphs, headings, and buttons. Without HTML, nothing exists!"),
        ("🎨 CSS (Cascading Style Sheets): ", "The Clothes, Paint & Super-Cape. It applies vibrant colors, glowing gradients, rounded corners, and trendy typography."),
        ("⚡ JavaScript (JS): ", "The Brain, Batteries & Muscles. It turns static documents into living video games! It animates character dances, calculates high scores, and shoots celebratory confetti.")
    ]
    for bold_prefix, desc in trio_items:
        p = doc.add_paragraph(style='List Bullet')
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.color.rgb = RGBColor(2, 132, 199)
        p.add_run(desc)

    doc.add_paragraph()

    # Section 4: Gamified Features
    h4 = doc.add_heading("4. Gamified & Interactive Features", level=1)
    h4.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    feats = [
        ("🚀 Packet Launch Mission: ", "Kids choose a destination (kitty-cafe.fun, space-arcade.io, pizza-bot.org) and press BLAST OFF to follow the rocket capsule in cinematic 3D tracking."),
        ("🏆 Kid's Web Quiz Quest: ", "4 playful multiple-choice questions with instant sound feedback, confetti bursts, and 4 collectible Cyber Badges (Junior Web Wizard)."),
        ("🔊 100% Procedural Audio Synth: ", "Built with Web Audio API. Generates bubble pops, rocket whooshes, calculation beeps, and fanfare chords with zero external audio downloads."),
        ("🗣️ Professor Byte Voice Narration: ", "Built-in SpeechSynthesis narrator reads station analogies and trivia in a lively, cheerful robotic mentor voice.")
    ]
    for b_prefix, desc in feats:
        p = doc.add_paragraph(style='List Bullet')
        r_b = p.add_run(b_prefix)
        r_b.bold = True
        p.add_run(desc)

    doc.add_paragraph()

    # Section 5: How to Run
    h5 = doc.add_heading("5. Quick Start & Execution Guide", level=1)
    h5.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    p5 = doc.add_paragraph(
        "To start the ATW WEB FLOW application:\n"
        "1. Double-click 'start.bat' in the project folder.\n"
        "   - It automatically checks port 5173.\n"
        "   - If any old server is already running, it terminates it.\n"
        "   - It automatically opens your web browser to http://localhost:5173/.\n"
        "   - Starts the web server cleanly with zero errors.\n\n"
        "Keyboard Controls:\n"
        "• Spacebar: Play / Pause 4D Time flow\n"
        "• Left / Right Arrow Keys: Scrub time backwards & forwards\n"
        "• Escape Key: Close any open dialog or modal"
    )

    # Save document
    out_file = os.path.abspath("ATW_Web_Flow_Guide.docx")
    doc.save(out_file)
    print(f"Word document successfully created: {out_file}")

if __name__ == "__main__":
    create_word_document()
